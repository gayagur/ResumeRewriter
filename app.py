
import streamlit as st
import requests
from docx import Document
from io import BytesIO

st.set_page_config(page_title="Smart Resume Rewriter", layout="wide")
st.title("🧠 Resume Rewriter")

API_KEY = "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"
MODEL = "mistralai/mistral-7b-instruct:free"
KEYWORDS = ["summary", "experience", "skills", "projects", "education"]

def rewrite_paragraph(paragraph_text, job_description, api_key):
    prompt = f"""You will receive a paragraph from a resume and a job description.

Please rewrite the paragraph so that it aligns better with the job description.
Do not change the meaning, just improve the wording to better fit the role.

Return only the updated paragraph, without explanations.

Resume paragraph:
{paragraph_text}

Job description:
{job_description}
"""
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json"
    }
    payload = {
        "model": MODEL,
        "messages": [
            {"role": "system", "content": "You are a professional resume editor."},
            {"role": "user", "content": prompt}
        ]
    }

    res = requests.post("https://openrouter.ai/api/v1/chat/completions", headers=headers, json=payload)
    if res.status_code != 200:
        raise Exception(f"Model error: {res.status_code} - {res.text}")
    return res.json()["choices"][0]["message"]["content"].strip()

def process_resume(doc, job_description, api_key):
    updated_doc = Document()
    rewritten_count = 0

    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            updated_doc.add_paragraph("")
            continue

        if any(kw in text.lower() for kw in KEYWORDS):
            updated_doc.add_paragraph(text, style=para.style)
            continue

        prev = updated_doc.paragraphs[-1].text.lower() if updated_doc.paragraphs else ""
        if any(kw in prev for kw in KEYWORDS):
            try:
                rewritten = rewrite_paragraph(text, job_description, api_key)
                new_p = updated_doc.add_paragraph(style=para.style)
                run = new_p.add_run(rewritten)
                if para.runs:
                    original_run = para.runs[0]
                    run.bold = original_run.bold
                    run.italic = original_run.italic
                    run.underline = original_run.underline
                    run.font.name = original_run.font.name
                    run.font.size = original_run.font.size
                rewritten_count += 1
            except Exception as e:
                updated_doc.add_paragraph(text, style=para.style)
        else:
            updated_doc.add_paragraph(text, style=para.style)

    buffer = BytesIO()
    updated_doc.save(buffer)
    buffer.seek(0)
    return buffer, rewritten_count

uploaded_file = st.file_uploader("Upload your resume (.docx)", type=["docx"])
job_description = st.text_area(" Paste job description", height=250)

if uploaded_file and job_description:
    if st.button(" Rewrite Resume✏️"):
        try:
            with st.spinner("Rewriting selected paragraphs..."):
                doc = Document(uploaded_file)
                updated_buffer, count = process_resume(doc, job_description, API_KEY)
                st.success(f"✅ Done! {count} paragraph(s) were updated.")
                st.download_button("⬇️ Download Updated Resume", updated_buffer, "updated_resume.docx",
                                   mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        except Exception as e:
            st.error(f"❌ Error: {e}")
